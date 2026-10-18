#!/usr/bin/env python3
"""Build a Soalin-format .docx from a questions JSON.

Usage:
  python build_soalin_docx.py --questions q_sourced.json --out bank.docx

Question object (see SKILL.md for the full schema):
  {"no": 1, "soal": "...", "gambar_soal": ["..."], "opsi": [{"huruf": "A", "teks": "...", "penjelasan": "..."}, ...],
   "kunci": "D", "gambar_penjelasan": "...", "gambar_penjelasan_slot": ["..."], "sumber": "...", "broken": false,
   "catatan_broken": "..."}

Rules enforced:
  - Numbering is kept exactly as given in "no" (no renumbering) unless --renumber is passed.
  - Each option line: "A. teks {penjelasan}". Missing penjelasan -> error (format requires it).
  - Image slots: soal images go right after the question text (before A). Penjelasan images go after
    "Gambar Penjelasan:" and before "Sumber:". Images are NEVER embedded, only [GAMBAR: ...] placeholders.
  - Broken questions (broken=true, or missing kunci) go to "Bagian 2" as a real Heading 2 + 3-column table.
  - Output is plain paragraphs (no bold/heading inside the question body).
"""
import argparse, json, sys
from docx import Document

def build(qs, out, renumber=False):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    broken_rows, slots = [], []
    n_out = 0
    for q in qs:
        no = q["no"]
        is_broken = q.get("broken") or not q.get("kunci")
        if is_broken:
            text = q["soal"] + " " + " ".join(f"{o['huruf']}. {o['teks']}" for o in q.get("opsi", []))
            if q.get("catatan_broken"):
                text += f" [{q['catatan_broken']}]"
            gcol = "[kosong — isi screenshot slide dosen]" if (q.get("gambar_soal") or q.get("gambar_penjelasan_slot")) else ""
            broken_rows.append((str(no), text, gcol))
            continue
        n_out += 1
        num = n_out if renumber else no
        doc.add_paragraph(f"{num}. {q['soal']}")
        for s in q.get("gambar_soal", []):
            doc.add_paragraph(f"[GAMBAR: {s}]"); slots.append(f"q{num} → gambar soal: {s}")
        for o in q["opsi"]:
            pj = o.get("penjelasan")
            if not pj:
                sys.exit(f"No {no} opsi {o['huruf']}: penjelasan opsi kosong (wajib)")
            doc.add_paragraph(f"{o['huruf']}. {o['teks']} {{{pj}}}")
        doc.add_paragraph(f"Kunci: {q['kunci']}")
        if q.get("gambar_penjelasan"):
            doc.add_paragraph(f"Gambar Penjelasan: {q['gambar_penjelasan']}")
        for s in q.get("gambar_penjelasan_slot", []):
            doc.add_paragraph(f"[GAMBAR: {s}]"); slots.append(f"q{num} → gambar penjelasan: {s}")
        if not q.get("sumber"):
            sys.exit(f"No {no}: sumber kosong (wajib)")
        doc.add_paragraph(f"Sumber: {q['sumber']}")
        doc.add_paragraph("")
    if broken_rows:
        doc.add_heading("Bagian 2", level=2)
        t = doc.add_table(rows=1, cols=3); t.style = "Table Grid"
        for c, h in zip(t.rows[0].cells, ["No", "Soal Asli (dari rekapan)", "Gambar Penjelasan"]):
            c.text = h
        for no, text, g in broken_rows:
            r = t.add_row().cells; r[0].text = no; r[1].text = text; r[2].text = g
    if slots:
        doc.add_paragraph("---CATATAN GAMBAR---")
        for s in slots:
            doc.add_paragraph(s)
    doc.save(out)
    print(f"saved {out}: {n_out} MCQ, {len(broken_rows)} broken, {len(slots)} image slots")

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--questions", required=True); ap.add_argument("--out", required=True)
    ap.add_argument("--renumber", action="store_true")
    a = ap.parse_args()
    build(json.load(open(a.questions, encoding="utf-8")), a.out, a.renumber)
